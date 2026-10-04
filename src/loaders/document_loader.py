from pathlib import Path

import pandas as pd
from pypdf import PdfReader
from docx import Document as DocxDocument

from langchain_core.documents import Document
from rich.console import Console


class DocumentLoader:

    def __init__(self, docs_path: str):

        self.docs_path = Path(docs_path)

        self.console = Console()

        self.pdf_count = 0
        self.docx_count = 0
        self.excel_count = 0

    def load(self):

        documents = []

        self.console.print(
            "\nStarting Document Loading...",
            style="bold green"
        )

        for file in self.docs_path.glob("*"):

            suffix = file.suffix.lower()

            try:

                if suffix == ".pdf":

                    docs = self._load_pdf(file)

                    documents.extend(docs)

                    self.pdf_count += 1

                elif suffix == ".docx":

                    docs = self._load_docx(file)

                    documents.extend(docs)

                    self.docx_count += 1

                elif suffix in [".xlsx", ".xls"]:

                    docs = self._load_excel(file)

                    documents.extend(docs)

                    self.excel_count += 1

                else:

                    self.console.print(
                        f"Skipping unsupported file: {file.name}",
                        style="yellow"
                    )

            except Exception as ex:

                self.console.print(
                    f"Error loading {file.name}",
                    style="bold red"
                )

                print(ex)

        self.print_summary(documents)

        return documents

    def _load_pdf(self, file_path: Path):

        self.console.print(
            f"Loading PDF : {file_path.name}",
            style="bold blue"
        )

        reader = PdfReader(str(file_path))

        documents = []

        for page_number, page in enumerate(
            reader.pages,
            start=1
        ):

            text = page.extract_text()

            if not text:
                continue

            documents.append(
                Document(
                    page_content=text,
                    metadata={
                        "source": file_path.name,
                        "file_type": "pdf",
                       "page": page_number
                    }
                )
            )

        return documents

    def _load_docx(self, file_path: Path):

        self.console.print(
            f"Loading DOCX : {file_path.name}",
            style="bold cyan"
        )

        doc = DocxDocument(str(file_path))

        text = "\n".join(
            paragraph.text
            for paragraph in doc.paragraphs
            if paragraph.text.strip()
        )

        return [
            Document(
                page_content=text,
                metadata={
                    "source": file_path.name,
                    "file_type": "docx"
                }
            )
        ]

    def _load_excel(self, file_path: Path):

        self.console.print(
            f"Loading XLSX : {file_path.name}",
            style="bold yellow"
        )

        workbook = pd.ExcelFile(file_path)

        documents = []

        for sheet_name in workbook.sheet_names:

            dataframe = pd.read_excel(
                file_path,
                sheet_name=sheet_name
            )

            text = dataframe.to_string(
                index=False
            )

            documents.append(
                Document(
                    page_content=text,
                    metadata={
                        "source": file_path.name,
                        "file_type": "xlsx",
                        "sheet_name": sheet_name
                    }
                )
            )

        return documents

    def print_summary(self, documents):

        print("\n" + "=" * 80)

        print("DOCUMENT LOADING SUMMARY")

        print("=" * 80)

        print(f"PDF Files Loaded   : {self.pdf_count}")
        print(f"DOCX Files Loaded  : {self.docx_count}")
        print(f"XLSX Files Loaded  : {self.excel_count}")
        print(f"Total Documents    : {len(documents)}")

        print("=" * 80)